"""Real format parsing and adversarial v0.2 controller behavior; no remote inference."""
from research_fixtures import create_test_raw
import asyncio
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
import json
from pathlib import Path
import tempfile
import time
import unittest
from unittest.mock import patch

from yantra.budget import BudgetExceeded, RunBudget, deadline
from yantra.engine import Engine, StopRun
from yantra.frameworks import FrameworkStore
from yantra.materials import numeric, parse_file, research_chunks
from yantra.policies import OfflineDemoBackend
from yantra.schemas import ChildArguments, MainAction, Result

PRICE = {"input_per_million": 2, "output_per_million": 8, "currency": "CNY", "source": "test-price", "as_of": "2026-10-13"}
TEXT = "2025年度 合并报表 CNY million 2026-01-01披露 Revenue 1,234.50 净利润 (123.45)\n公司表示面临行业竞争风险，产能扩张存在不确定性。"


class V2TestBackend(OfflineDemoBackend):
    last_usage = None

    async def choose_framework(self, payload):
        return {"identifier": "quality", "reason": "现金流口径尚需核对"}

    async def framework(self, payload):
        return {"id":"custom-quality","name":"用户框架","scope":"一般公司现金流质量","questions":["收入与现金转化"],"evidence_needs":["原始财报"],"risks":["竞争"],"disconfirming":["现金转化恶化"],"tracking":["经营现金流"]}

    async def review_framework(self, payload):
        return {"passed":True,"reasons":["范围、证据与反证完整"]}

    async def extract(self, payload):
        doc = payload["document"]
        anchor = doc["anchors"][0]
        return {"financials": [{"field": "revenue", "value": 1234.5, "raw_value": "1,234.50", "period": "FY2025", "unit": "million", "currency": "CNY", "consolidated": True,
            "disclosure_date": "2026-01-01", "source_ref": anchor["ref"], "source_hash": doc["source_hash"]}],
            "evidence": [{"statement": "披露收入", "kind": "fact", "quote": "Revenue 1,234.50", "source_ref": anchor["ref"], "source_hash": doc["source_hash"], "disclosure_date": "2026-01-01", "question": "收入基础"}], "gaps": ["未取得FCFF调节与行情"]}

    async def main(self, view):
        a = view["artifacts"]
        if not a.get("finance"):
            return self._main("delegate_task", target="finance", purpose="collect")
        if not a.get("thesis"):
            return self._main("delegate_task", target="thesis", purpose="analyze")
        if not a.get("proposal"):
            return self._main("delegate_task", target="valuation", purpose="propose")
        if not a.get("verification"):
            return self._main("delegate_task", target="verification", purpose="audit")
        return self._main("finalize")

    async def worker(self, task, local):
        outputs = local["tool_outputs"]
        successful = {o["tool"]: o["response"] for o in outputs if o["response"]["ok"]}
        if task["depth"] == 3:
            if not successful:
                return self._tool("read_materials", {"company": task["company"]})
            return self._result(task, local, data={"checked": task["dynamic_role"]})
        if task["target"] == "finance":
            if "read_materials" not in successful:
                return self._tool("read_materials", {"company": task["company"]})
            children = len([x for x in outputs if x["tool"] == "spawn_subtask"])
            if children < 2:
                return self._tool("spawn_subtask", {"role": ("披露日期核对", "经营风险定位")[children], "question": "核对指定材料", "tools": ["read_materials"], "completion_conditions": ["返回原文引用"]})
            if "extract_financials" not in successful:
                return self._tool("extract_financials", {"company": task["company"]})
            extracted = successful["extract_financials"]["data"]
            if "calculate_metrics" not in successful:
                return self._tool("calculate_metrics", {"records": extracted["financials"]})
            return self._result(task, local, data={"use_receipts": True}, unresolved=extracted["gaps"])
        if task["target"] == "thesis":
            refs = task["inputs"]["finance"]["evidence_refs"]
            return self._result(task, local, data={"claims": [{"type": "fact", "statement": "收入披露可定位", "evidence_refs": refs}], "hypotheses": [{"type": "assumption", "statement": "盈利改善待核验", "evidence_refs": refs}],
                "summary": "已取得收入证据，FCFF及价格不足，保留条件性论点。", "risks": ["行业竞争"], "disconfirming": ["收入下滑"], "tracking": ["现金流"], "expectations": {"status": "unverified_hypothesis", "baseline": None, "statement": "缺市场基准"}, "gaps": ["缺行情"]})
        if task["target"] == "valuation":
            return self._result(task, local, status="partial", data={"model": "not_applicable", "parameters": {}, "basis": {}, "scope": "FCFF调节、股数及稳定性不足", "gaps": ["不能可靠DCF"]})
        if "check_consistency" not in successful:
            return self._tool("check_consistency", {"results": task["inputs"], "assumptions": task["assumptions"]})
        return self._result(task, local, data=successful["check_consistency"]["data"])


class PositiveCashFlowBackend(V2TestBackend):
    async def extract(self, payload):
        doc = payload["document"]
        anchor = doc["anchors"][0]
        financials = []
        for field, value in {"revenue": 1234.5, "operating_cash_flow": 192, "capex": 72, "debt": 150, "cash": 50, "shares": 100, "minority_interest": 0, "ebit": 160, "tax_rate": .25, "depreciation": 72, "delta_nwc": 0}.items():
            financials.append({"field": field, "value": value, "raw_value": "1,234.50" if field=="revenue" else str(value), "period": "FY2025", "unit": "ratio" if field=="tax_rate" else "million shares" if field=="shares" else "million", "currency": "CNY", "consolidated": True,
                "source_ref": anchor["ref"], "source_hash": doc["source_hash"], "disclosure_date": "2026-01-01",
                "cash_flow_basis": "unlevered" if field=="operating_cash_flow" else None, "security_code": "600000", "share_class": "A", "accounting_scope": "consolidated", "share_scope": "total ordinary shares" if field=="shares" else None})
        value = await super().extract(payload)
        value["financials"] = financials
        return value

    async def main(self, view):
        if view["artifacts"].get("proposal") and view["artifacts"].get("thesis"):
            if not view["assumptions"]["confirmed"]:
                return self._main("ask_user", question="确认当前假设")
            if not view["artifacts"].get("valuation"):
                return self._main("delegate_task", target="valuation", purpose="calculate")
        return await super().main(view)

    async def worker(self, task, local):
        if task["target"] == "valuation":
            if task["purpose"] == "propose":
                parameters = {"base_fcf":120,"growth_rate":.05,"discount_rate":.1,"terminal_growth":.02,"years":5,"net_debt":100,"shares":100,"minority_interest":0,"currency":"CNY","unit":"million"}
                refs = task["inputs"]["finance"]["evidence_refs"]
                return self._result(task, local, data={"model":"dcf","parameters":parameters,"basis":{k:{"source_refs":refs,"formula":"固定原始材料及显式测试假设"} for k in parameters},"scope":"虚构稳定正无杠杆现金流工业样本","gaps":[]})
            if not local["tool_outputs"]:
                return self._tool("valuation", {"parameters":task["assumptions"]["parameters"],"assumptions_version":task["assumptions"]["version"]})
            return self._result(task, local, data=local["tool_outputs"][-1]["response"]["data"])
        return await super().worker(task, local)


def fake_search(adapter, query, purpose, max_results):
    return {"query": query, "purpose": purpose, "results": [], "errors": ["test_no_network"], "acquired_at": "2026-10-13T00:00:00Z", "service_cost_cny": 0}


class BudgetTests(unittest.TestCase):
    def make_budget(self, directory, cap=1):
        state = {"deadline": datetime.fromtimestamp(time.time()+100, timezone.utc).isoformat(), "limits": {"close_seconds": 1, "max_cost_cny": cap}, "pricing": PRICE, "model_ledger": [], "budget": {}, "stop_reasons": []}
        budget = RunBudget(state, lambda: None)
        budget.state_path = directory
        return budget

    def test_atomic_reservation_and_missing_usage(self):
        with tempfile.TemporaryDirectory() as d:
            budget = self.make_budget(Path(d), .05)
            def reserve(_):
                try:
                    return budget.reserve("x", "instruction", 4096)
                except BudgetExceeded:
                    return None
            with ThreadPoolExecutor(max_workers=2) as pool:
                receipts = list(pool.map(reserve, [1, 2]))
            self.assertEqual(sum(x is not None for x in receipts), 1)
            receipt = next(x for x in receipts if x)
            budget.settle(receipt, None)
            self.assertGreater(budget.state["budget"]["model_cost_cny"], 0)
            self.assertNotIn("cost_cny", receipt)

    def test_deadline_wait_and_cancel(self):
        with tempfile.TemporaryDirectory() as d:
            budget = self.make_budget(Path(d))
            budget.clock = lambda: time.time()+200
            self.assertRaises(BudgetExceeded, budget.check)
            budget.clock = time.time
            (Path(d)/"cancel.request").touch()
            self.assertRaises(BudgetExceeded, budget.check)

    def test_two_hour_deadline_with_closing_reserve(self):
        with tempfile.TemporaryDirectory() as d:
            budget = self.make_budget(Path(d))
            budget.state["deadline"] = deadline(0, 7200)
            budget.state["limits"]["close_seconds"] = 30
            budget.clock = lambda: 7169
            budget.check()
            budget.clock = lambda: 7170
            self.assertRaises(BudgetExceeded, budget.check)
            self.assertEqual(budget.remaining(closing=True), 30)
            self.assertRaises(ValueError, deadline, 0, 7200.01)

    def test_usage_settlement(self):
        with tempfile.TemporaryDirectory() as d:
            budget = self.make_budget(Path(d))
            receipt = budget.reserve("123", "test", 100)
            budget.settle(receipt, {"input_tokens": 100, "output_tokens": 10})
            self.assertAlmostEqual(receipt["cost_cny"], .00028)


class FormatTests(unittest.TestCase):
    def test_financial_heading_selection_and_foreign_language_query(self):
        chunks=[{"ref":"a1","text":"Annual report cover","location":{"page":1}},
                {"ref":"a2","text":"Contents: consolidated statement of cash flows 278","location":{"page":2}},
                {"ref":"a3","text":"CONSOLIDATED STATEMENT OF\nCASH FLOWS\nOperating cash flow 1,234.50 2,345.60 3,456.70","location":{"page":278}},
                {"ref":"a4","text":"Notes to the consolidated financial statements\n"+"other disclosure "*40+"consolidated statement of cash flows","location":{"page":310}}]
        selected=research_chunks({"chunks":chunks},query="合并现金流量表")
        self.assertTrue(any(c['ref']=='a3' for c in selected))
        fallback=research_chunks({"chunks":chunks},query="不存在的主题")
        self.assertTrue(any(c['ref']=='a3' for c in fallback))
    def test_office_web_and_structured_anchors(self):
        from docx import Document
        from pptx import Presentation
        from pptx.util import Inches
        import openpyxl
        with tempfile.TemporaryDirectory() as d:
            p = Path(d)
            doc = Document(); doc.add_paragraph("营业收入 1,234.50"); doc.add_table(rows=1, cols=2).cell(0, 0).text = "净利润 (123.45)"; doc.save(p/"sample.docx")
            ppt = Presentation(); slide = ppt.slides.add_slide(ppt.slide_layouts[6]); slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = "Revenue 1,234.50"; ppt.save(p/"sample.pptx")
            book = openpyxl.Workbook(); book.active.append(["营业收入", 1234.5]); book.save(p/"sample.xlsx")
            (p/"sample.html").write_text("<html><main><p>收入 1,234.50</p></main><script>fake 0</script></html>")
            (p/"sample.csv").write_text("field,value\nrevenue,1234.5\n")
            (p/"sample.json").write_text('{"revenue":1234.5}')
            for file in p.iterdir():
                chunks, gaps = parse_file(file)
                self.assertTrue(chunks, file.name)
                self.assertTrue(any("1234" in c["text"] or "1,234" in c["text"] for c in chunks), file.name)
                self.assertTrue(all(c["location"] for c in chunks))
            self.assertEqual(numeric("(1,234.50)"), -1234.5)

    def test_text_pdf_and_actual_english_chinese_ocr(self):
        from PIL import Image, ImageDraw, ImageFont
        from reportlab.pdfgen.canvas import Canvas
        with tempfile.TemporaryDirectory() as d:
            p = Path(d)
            canvas = Canvas(str(p/"native.pdf")); canvas.drawString(72, 720, "Revenue 1,234.50 million CNY"); canvas.save()
            chunks, _ = parse_file(p/"native.pdf")
            self.assertIn("1,234.50", chunks[0]["text"])
            self.assertEqual(chunks[0]["location"]["page"], 1)
            image = Image.new("RGB", (1600, 350), "white"); draw = ImageDraw.Draw(image)
            font_path = Path("/System/Library/Fonts/STHeiti Medium.ttc")
            if not font_path.is_file():
                font_path = Path("/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc")
            font = ImageFont.truetype(str(font_path), 56)
            draw.text((30, 30), "营业收入 1234.50 万元", font=font, fill="black")
            draw.text((30, 140), "Revenue 1234.50 CNY", font=font, fill="black")
            image.save(p/"scan.png"); image.save(p/"scan.pdf", "PDF")
            for file in (p/"scan.png", p/"scan.pdf"):
                chunks, gaps = parse_file(file)
                self.assertTrue(any("1234.50" in x["text"] for x in chunks), chunks)
                self.assertTrue(any("营业" in x["text"] for x in chunks), chunks)
                self.assertTrue(all("region_pixels" in x["location"] for x in chunks))


class ControllerTests(unittest.IsolatedAsyncioTestCase):
    def create(self, d, **kwargs):
        source = Path(d)/"original.txt"; source.write_text(TEXT)
        return create_test_raw(Path(d)/"runs", {"company": "测试公司", "horizon": "12个月", "cutoff": "2026-10-13T04:30:00+08:00", "market": "CN", "security_code": "600000", "trading_currency": "CNY", "confirmation_mode": "autonomous"},
            [source], backend=V2TestBackend(), pricing=PRICE, **kwargs)

    async def test_raw_flow_dynamic_roles_and_no_forced_valuation(self):
        from yantra.adk_app import run_engine
        with tempfile.TemporaryDirectory() as d, patch("yantra.materials.RawMaterialAdapter.search", fake_search):
            engine = self.create(d)
            state = await run_engine(engine)
            self.assertEqual(state["status"], "partial", state["stop_reasons"])
            self.assertEqual({q["purpose"] for q in state["searches"]}, {"disclosure", "counter_evidence", "changes", "expectations"})
            self.assertEqual(len([t for t in state["tasks"] if t["depth"] == 3]), 2)
            self.assertTrue(state["artifacts"]["finance"]["data"]["financials"])
            self.assertTrue(state["artifacts"]["finance"]["data"]["evidence"])
            self.assertTrue(state["artifacts"]["verification"])
            loaded = Engine.load(Path(d)/"runs", state["run_id"])
            self.assertEqual(loaded.state["deadline"], state["deadline"])
            self.assertEqual(loaded.state["model_ledger"], state["model_ledger"])
            self.assertTrue(Path(state["report_path"]).is_file())

    async def test_depth_and_tools_cannot_expand(self):
        with tempfile.TemporaryDirectory() as d:
            engine = self.create(d)
            self.assertEqual(engine.state["limits"]["run_seconds"], 7200)
            self.assertEqual(datetime.fromisoformat(engine.state["deadline"]).timestamp()-datetime.fromisoformat(engine.state["created_at"]).timestamp(),7200)
            task = engine._make_task(MainAction(action="delegate_task", target="finance", purpose="collect"))
            spec = {"role": "任意角色", "question": "test", "tools": ["valuation"], "completion_conditions": ["done"]}
            with self.assertRaises(StopRun):
                await engine._spawn(task, spec)
            task.depth = 3
            with self.assertRaises(StopRun):
                await engine._spawn(task, spec | {"tools": ["read_materials"]})

    async def test_verifier_veto_cannot_be_overridden(self):
        with tempfile.TemporaryDirectory() as d:
            engine = self.create(d)
            task = engine._make_task(MainAction(action="delegate_task", target="verification", purpose="audit"))
            result = Result(task_id=task.task_id, status="partial", data={"passed": False, "issues": []}, assumptions_version=1, actual_calls=1)
            local = {"actual_calls": 1, "tool_outputs": [{"tool": "check_consistency", "response": {"ok": True, "data": {"passed": True, "issues": []}, "evidence_refs": [], "calculation_refs": []}}]}
            self.assertFalse(engine._validate_worker_result(task, local, result)["data"]["passed"])

    async def test_future_or_fabricated_value_rejected(self):
        with tempfile.TemporaryDirectory() as d:
            engine = self.create(d)
            doc = engine.tools.adapter.read_materials(__import__('yantra.schemas',fromlist=['ReadArguments']).ReadArguments(company="测试公司"))["documents"][0]
            extraction = await engine.backend.extract({"document": {"source_hash": doc["source_hash"], "anchors": doc["chunks"]}})
            extraction["financials"][0]["value"] = 999
            accepted = engine.tools.accept_extraction(extraction, engine.state["research"]["cutoff"])
            self.assertEqual(accepted["financials"], [])

    async def test_derived_base_unit_value_is_not_double_scaled(self):
        with tempfile.TemporaryDirectory() as d:
            engine=self.create(d)
            doc=engine.tools.adapter.read_materials(__import__('yantra.schemas',fromlist=['ReadArguments']).ReadArguments(company="测试公司"))["documents"][0]
            extraction=await engine.backend.extract({"document":{"source_hash":doc['source_hash'],"anchors":doc['chunks']}})
            row=extraction['financials'][0]
            row['unit']='thousand';row['value']=1234500
            accepted=engine.tools.accept_extraction(extraction,engine.state['research']['cutoff'])
            self.assertEqual(accepted['financials'][0]['value'],1234.5)
            self.assertEqual(accepted['financials'][0]['model_value'],1234500)
            self.assertTrue(accepted['financials'][0]['normalization_note'])
            extraction["financials"][0]["value"] = 1234.5
            extraction["financials"][0]["disclosure_date"] = "2027-01-01"
            accepted = engine.tools.accept_extraction(extraction, engine.state["research"]["cutoff"])
            self.assertEqual(accepted["financials"], [])

    async def test_shared_read_message_and_approved_transfer(self):
        with tempfile.TemporaryDirectory() as d:
            engine = self.create(d)
            finance = engine._make_task(MainAction(action="delegate_task", target="finance", purpose="collect"))
            result = Result(task_id=finance.task_id, status="partial", data={"original": "shared"}, assumptions_version=1, actual_calls=0).model_dump()
            engine.state["tasks"].append(finance.model_dump()); engine.state["results"].append(result)
            engine.state["artifacts"]["finance"] = result
            thesis = engine._make_task(MainAction(action="delegate_task", target="thesis", purpose="analyze"))
            engine.state["tasks"].append(thesis.model_dump())
            read = engine._collaborate(thesis, {"operation": "read", "target_task_id": finance.task_id})
            self.assertEqual(read["data"]["data"]["original"], "shared")
            engine._collaborate(thesis, {"operation": "message", "target_task_id": finance.task_id, "message": "发现口径问题"})
            self.assertEqual(len(engine.state["tasks"]), 2)
            with self.assertRaises(StopRun):
                engine._collaborate(finance, {"operation": "read", "target_task_id": thesis.task_id})
            request = engine._collaborate(thesis, {"operation": "request_transfer", "target_task_id": finance.task_id,
                "message": "补充核对日期", "completion_conditions": ["返回有出处日期"], "tool_budget": 1})["data"]
            self.assertEqual(len(engine.state["tasks"]), 2)
            async def work(task):
                return Result(task_id=task["task_id"], status="partial", assumptions_version=1, actual_calls=0).model_dump()
            with patch.object(engine, "_work_task", side_effect=work):
                await engine._approve_transfer(MainAction(action="approve_transfer", transfer_id=request["id"], rationale="核对日期属于原任务"))
            self.assertEqual(engine.state["transfers"][0]["status"], "executed")
            self.assertEqual(len(engine.state["tasks"]), 3)

    async def test_manual_wait_expires_without_approval_or_budget_reset(self):
        with tempfile.TemporaryDirectory() as d:
            engine = self.create(d)
            engine.state.update(status="waiting", prepared=True)
            engine.state["research"]["confirmation_mode"] = "manual"
            engine.state["deadline"] = "2020-01-01T00:00:00+00:00"
            engine._save()
            state = await engine.advance()
            self.assertNotEqual(state["status"], "waiting")
            self.assertFalse(state["assumptions"]["confirmed"])
            self.assertEqual(state["budget"]["model_calls"], 0)
            self.assertEqual(Engine.load(Path(d)/"runs",state["run_id"]).state["deadline"], "2020-01-01T00:00:00+00:00")

    async def test_raw_replay_recomputes_without_remote_calls(self):
        from yantra.adk_app import run_engine
        from yantra.replay import replay_raw
        with tempfile.TemporaryDirectory() as d, patch("yantra.materials.RawMaterialAdapter.search", fake_search):
            engine = self.create(d)
            state = await run_engine(engine)
            replayed, comparison = await replay_raw(Path(d)/"runs", state)
            self.assertTrue(comparison["calculation_equal"])
            self.assertEqual(comparison["remote_model_calls"], 0)
            self.assertEqual(replayed.state["framework"], state["framework"])

    async def test_failed_supplement_retains_validated_records(self):
        from yantra.adk_app import run_engine
        with tempfile.TemporaryDirectory() as d, patch("yantra.materials.RawMaterialAdapter.search", fake_search):
            engine = self.create(d)
            await run_engine(engine)
            task = engine._make_task(MainAction(action="delegate_task",target="finance",purpose="supplement"))
            partial = engine._partial_worker_result(task,{"tool_outputs":[],"actual_calls":0},"source unavailable")
            self.assertTrue(partial["data"]["financials"])
            self.assertTrue(partial["data"]["evidence"])
            self.assertIn("source unavailable",partial["unresolved"])

    async def test_user_framework_is_verified_and_used_in_current_run(self):
        from yantra.adk_app import run_engine
        with tempfile.TemporaryDirectory() as d, patch("yantra.materials.RawMaterialAdapter.search",fake_search):
            path=Path(d)/"framework.md";path.write_text("关注收入与现金转化，查原始财报，以现金转化恶化作为反证。")
            engine=self.create(d,framework_file=path)
            state=await run_engine(engine)
            self.assertEqual(state['framework']['rules']['id'],'custom-quality')
            self.assertTrue(state['framework']['verified'])
            self.assertEqual(state['framework']['original'],path.read_text())
            self.assertTrue(any(t['target']=='verification' for t in state['tasks']))

    async def test_autonomous_and_manual_authorization_then_modify_recomputes(self):
        from yantra.adk_app import run_engine
        for mode in ("autonomous", "manual"):
            with tempfile.TemporaryDirectory() as d, patch("yantra.materials.RawMaterialAdapter.search", fake_search):
                source=Path(d)/"original.txt"
                source.write_text(TEXT.splitlines()[0]+" unlevered operating_cash_flow 192 capex 72 debt 150 cash 50 total ordinary shares 100 minority_interest 0 ebit 160 tax_rate 0.25 depreciation 72 delta_nwc 0\n")
                research={"company":"测试公司","horizon":"12个月","cutoff":"2026-10-13","market":"CN","security_code":"600000","trading_currency":"CNY","confirmation_mode":mode}
                engine=create_test_raw(Path(d)/"runs",research,[source],backend=PositiveCashFlowBackend(),pricing=PRICE,framework="quality")
                state=await run_engine(engine)
                if mode=="manual":
                    self.assertEqual(state["status"],"partial",state["stop_reasons"])
                    self.assertFalse(state["assumptions"]["confirmed"])
                    self.assertEqual(state["assumptions"]["authorization_source"], "provisional")
                    state=await engine.respond("confirm")
                self.assertIn("valuation",state["artifacts"],state["stop_reasons"])
                self.assertTrue(state["artifacts"]["verification"]["data"]["passed"],state["issues"])
                self.assertEqual(state["assumptions"]["authorization_source"], "controller" if mode=="autonomous" else "user")
                old=state["artifacts"]["valuation"]["data"]["per_share_value"]
                changed=await engine.respond("modify",{"growth_rate":.03})
                self.assertEqual(changed["assumptions"]["version"],2 if mode=="autonomous" else 3)
                self.assertNotEqual(old,changed["artifacts"]["valuation"]["data"]["per_share_value"])
                self.assertTrue(changed["invalidated"])


class FrameworkTests(unittest.TestCase):
    def test_draft_and_pinned_baseline(self):
        with tempfile.TemporaryDirectory() as d:
            store = FrameworkStore(d)
            baseline = store.builtin("quality")
            draft = store.draft(baseline["rules"], "原文框架", baseline["version"], "调整", "run-test")
            self.assertRaises(ValueError, store.load, draft["version"])
            store.verify(draft["version"], {"passed": True, "review_task_id": "verification-task"})
            self.assertTrue(store.load(draft["version"])["verified"])
            self.assertIsNone(baseline.get("verification"))
