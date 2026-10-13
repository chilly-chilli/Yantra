"""Raw bytes, format-specific anchors, OCR and public-source snapshots."""
from __future__ import annotations

import csv
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re
from urllib.parse import urlparse

import httpx

from .schemas import ReadArguments

FORMATS = {".pdf", ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".xlsx", ".docx", ".pptx", ".txt", ".md", ".html", ".htm", ".csv", ".json", ".xml"}


def research_chunks(doc, limit=12, query=None, anchor_refs=None):
    chunks = doc["chunks"]
    if anchor_refs is not None:
        return [c for c in chunks if c["ref"] in anchor_refs][:limit]
    if query:
        words = re.findall(r"[A-Za-z]+|[\u4e00-\u9fff]+", query.lower())
        translations = {"现金流": "cash flow", "资产负债": "balance sheet", "合并利润": "consolidated income", "营业收入": "revenue", "股本": "share capital", "股数": "shares", "债务": "borrowings", "资本开支": "capital expenditure"}
        words.extend(v for k, v in translations.items() if k in query)
        def relevance(chunk):
            text = re.sub(r"\s+", " ", chunk["text"].lower())
            matches = sum(text.count(w) for w in words)
            return matches * 10 + min(20, len(re.findall(r"\d,\d{3}|\d\.\d", chunk["text"]))) if matches else 0
        if any(relevance(c) for c in chunks):
            return sorted(chunks, key=relevance, reverse=True)[:limit]
        return research_chunks(doc, limit)
    groups = [(("consolidated income statement", "consolidated statement of profit", "合并利润表", "净利润"), 2),
              (("consolidated statement of cash flows", "consolidated cash flow", "合并现金流量表", "经营活动产生的现金流量净额"), 2),
              (("consolidated balance sheet", "consolidated statement of financial position", "合并资产负债表"), 2),
              (("revenue", "营业收入"), 2), (("risk", "competition", "风险", "竞争"), 3)]
    def score(chunk, words):
        text = re.sub(r"\s+", " ", chunk["text"].lower())
        matches = sum(text.count(word) * (30 if "consolidated" in word or "合并" in word else 1) for word in words)
        heading = 200 if any(word in text[:200] and ("consolidated" in word or "合并" in word) for word in words) else 0
        return heading + matches + min(20, len(re.findall(r"\d,\d{3}|\d\.\d", text))) if matches else 0
    selected = chunks[:1]
    for words, count in groups:
        selected.extend(sorted(chunks, key=lambda c: score(c, words), reverse=True)[:count])
    seen = set()
    output = []
    for chunk in selected:
        if chunk["ref"] not in seen:
            seen.add(chunk["ref"])
            output.append(chunk)
    return output[:limit]


def now():
    return datetime.now(timezone.utc).isoformat()


def cutoff_time(value):
    parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
    if parsed.tzinfo is None:
        from zoneinfo import ZoneInfo
        parsed = parsed.replace(tzinfo=ZoneInfo("Asia/Shanghai"))
        if len(value) == 10:
            parsed = parsed.replace(hour=23, minute=59, second=59, microsecond=999999)
    return parsed


def numeric(value):
    value = value.strip().replace(",", "").replace("，", "").replace("−", "-").replace(" ", "")
    if value.startswith(("(", "（")) and value.endswith((")", "）")):
        value = "-" + value[1:-1]
    if not re.fullmatch(r"[-+]?\d+(?:\.\d+)?", value):
        raise ValueError("not an exact numeric source value")
    return float(value)


def source_number_present(raw, text):
    """Require a complete original numeric token, not a digit inside a code/date/value."""
    return bool(raw and re.search(r"(?<![\d.,+\-(（])" + re.escape(raw) + r"(?![\d.,)）])", text))


def parse_file(path, check=lambda: None, timeout=30, recognizer=None):
    """Each location points to the original format, never only converted Markdown."""
    path = Path(path)
    suffix = path.suffix.lower()
    chunks, gaps = [], []

    def add(text, location):
        if str(text).strip():
            chunks.append({"text": str(text), "location": location})

    def ocr(image, location):
        if recognizer is not None:
            result=recognizer(image,timeout)
            for row in result["words"]:
                add(row["text"],{**location,"region_pixels":row["bbox"],"ocr":True})
            gaps.extend(result.get("warnings",[]))
            gaps.append("OCR 表格列关系及精确数字需原图核对；图表数值未推断")
            return
        import pytesseract
        import os
        import sys
        bundled = Path(sys.prefix) / "share" / "tessdata"
        if (bundled / "chi_sim.traineddata").is_file():
            os.environ.setdefault("TESSDATA_PREFIX", str(bundled))
        languages = set(pytesseract.get_languages())
        lang = "chi_sim+eng" if "chi_sim" in languages else "eng"
        if "chi_sim" not in languages:
            gaps.append("未安装 chi_sim，中文 OCR 能力未满足")
        data = pytesseract.image_to_data(image, lang=lang, output_type=pytesseract.Output.DICT, timeout=timeout)
        groups = {}
        for i, word in enumerate(data["text"]):
            if not word.strip():
                continue
            key = (data["block_num"][i], data["par_num"][i], data["line_num"][i])
            groups.setdefault(key, []).append(i)
        for indices in groups.values():
            left = min(data["left"][i] for i in indices)
            top = min(data["top"][i] for i in indices)
            right = max(data["left"][i] + data["width"][i] for i in indices)
            bottom = max(data["top"][i] + data["height"][i] for i in indices)
            add(" ".join(data["text"][i] for i in indices), {**location, "region_pixels": [left, top, right, bottom], "ocr": True})
        gaps.append("OCR 表格列关系及精确数字需原图核对；图表数值未推断")

    check()
    if suffix == ".pdf":
        import pdfplumber
        import pypdfium2
        with pdfplumber.open(path) as document:
            for page_number, page in enumerate(document.pages, 1):
                check()
                text = page.extract_text() or ""
                if len(text.strip()) >= 20:
                    add(text, {"page": page_number, "bbox": list(page.bbox)})
                    for i, table in enumerate(page.find_tables(), 1):
                        add("\n".join(" | ".join(str(c or "") for c in row) for row in table.extract()),
                            {"page": page_number, "table": i, "bbox": list(table.bbox)})
                else:
                    rendered = pypdfium2.PdfDocument(path)
                    try:
                        ocr(rendered[page_number - 1].render(scale=3).to_pil(), {"page": page_number, "render_scale": 3})
                    finally:
                        rendered.close()
        gaps.append("PDF 图表、无文本区域及复杂表格仍需视觉核对")
    elif suffix in {".png", ".jpg", ".jpeg", ".tif", ".tiff"}:
        from PIL import Image, ImageSequence
        with Image.open(path) as image:
            for frame, item in enumerate(ImageSequence.Iterator(image), 1):
                check()
                ocr(item.convert("RGB"), {"image": path.name, "frame": frame})
    elif suffix == ".xlsx":
        import openpyxl
        book = openpyxl.load_workbook(path, data_only=False, read_only=True)
        cached = openpyxl.load_workbook(path, data_only=True, read_only=True)
        try:
            for sheet in book:
                for row in sheet:
                    check()
                    cells = []
                    for cell in row:
                        if cell.value is not None:
                            value = cell.value
                            if cell.data_type == "f":
                                value = cached[sheet.title][cell.coordinate].value
                                if value is None:
                                    gaps.append(f"Excel 公式缺少缓存值：{sheet.title}!{cell.coordinate}")
                                    continue
                            cells.append(f"{cell.coordinate}={value}")
                    if cells:
                        add(" | ".join(cells), {"sheet": sheet.title, "cells": [c.coordinate for c in row if c.value is not None]})
        finally:
            book.close()
            cached.close()
    elif suffix == ".docx":
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        for index, block in enumerate(Document(path).iter_inner_content(), 1):
            check()
            if isinstance(block, Paragraph):
                add(block.text, {"block": index, "paragraph": index})
            elif isinstance(block, Table):
                for row_number, row in enumerate(block.rows, 1):
                    add(" | ".join(cell.text for cell in row.cells), {"block": index, "table_row": row_number})
        gaps.append("Word 嵌入图片／图表未自动提取，需另交原图")
    elif suffix == ".pptx":
        from pptx import Presentation
        for slide_number, slide in enumerate(Presentation(path).slides, 1):
            check()
            for shape_number, shape in enumerate(slide.shapes, 1):
                if shape.has_text_frame:
                    add(shape.text, {"slide": slide_number, "shape": shape_number})
                if shape.has_table:
                    for row_number, row in enumerate(shape.table.rows, 1):
                        add(" | ".join(cell.text for cell in row.cells), {"slide": slide_number, "shape": shape_number, "table_row": row_number})
        gaps.append("PPT 图片及图表数字需原图核对")
    elif suffix in {".html", ".htm"}:
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(path.read_bytes(), "html.parser")
        for node in soup(["script", "style", "nav", "footer"]):
            node.decompose()
        article = soup.find("main") or soup.find("article") or soup.body or soup
        for index, line in enumerate(article.get_text("\n", strip=True).splitlines(), 1):
            add(line, {"web_fragment": index})
    elif suffix == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as file:
            for row_number, row in enumerate(csv.reader(file), 1):
                add(" | ".join(row), {"csv_row": row_number})
    elif suffix == ".json":
        data = json.loads(path.read_text(encoding="utf-8-sig"))
        def walk(value, pointer=""):
            if isinstance(value, dict):
                for k, v in value.items():
                    walk(v, pointer + "/" + str(k).replace("~", "~0").replace("/", "~1"))
            elif isinstance(value, list):
                for i, v in enumerate(value):
                    walk(v, pointer + "/" + str(i))
            else:
                add(f"{pointer}: {value}", {"json_pointer": pointer})
        walk(data)
    elif suffix in {".txt", ".md", ".xml"}:
        for i, line in enumerate(path.read_text(encoding="utf-8-sig").splitlines(), 1):
            add(line, {"line": i})
    else:
        raise ValueError(f"unsupported material format: {suffix}")
    if not chunks:
        gaps.append("未提取到正文或表格")
    return chunks, list(dict.fromkeys(gaps))


class RawMaterialAdapter:
    def __init__(self, directory):
        self.materials_dir = Path(directory)
        self.check = lambda: None
        self.timeout = lambda: 45
        self.recognizer = None
        self.cached_parse = None

    @classmethod
    def create(cls, directory, company, inputs, check=lambda: None, recognizer=None, cached_parse=None):
        self = cls(directory)
        directory = self.materials_dir
        directory.mkdir(parents=True, exist_ok=True)
        (directory / "index.json").write_text(json.dumps({"company": company, "company_name": company, "version": 1,
            "adapter": "raw-v2", "fictional": False, "materials": [], "gaps": []}, ensure_ascii=False))
        self.check = check
        self.recognizer = recognizer
        self.cached_parse = cached_parse
        for item in inputs:
            if str(item).startswith(("https://", "http://")):
                try:
                    self.fetch(str(item))
                except Exception as exc:
                    self.gap(f"获取输入链接失败 {item}: {type(exc).__name__}")
            else:
                path = Path(item).expanduser().resolve()
                if not path.exists():
                    raise FileNotFoundError(str(path))
                paths = sorted(p for p in path.rglob("*") if p.is_file() and p.suffix.lower() in FORMATS) if path.is_dir() else [path]
                for file in paths:
                    check()
                    self.add(file.read_bytes(), file.suffix.lower(), str(file))
        return self

    def load_manifest(self):
        return json.loads((self.materials_dir / "index.json").read_text())

    def _write(self, index):
        temporary = self.materials_dir / "index.tmp"
        temporary.write_text(json.dumps(index, ensure_ascii=False, indent=2))
        temporary.replace(self.materials_dir / "index.json")

    def gap(self, message):
        index = self.load_manifest()
        if message not in index["gaps"]:
            index["gaps"].append(message)
        self._write(index)

    def add(self, content, suffix, origin, url=None, disclosure_date=None, channel=None):
        self.check()
        digest = hashlib.sha256(content).hexdigest()
        index = self.load_manifest()
        existing = next((x for x in index["materials"] if x["source_hash"] == digest), None)
        if existing:
            existing.setdefault("acquisitions", []).append(channel or {"origin": origin, "url": url, "acquired_at": now()})
            self._write(index)
            return existing
        identifier = "src-" + digest[:20]
        path = self.materials_dir / (identifier + suffix)
        path.write_bytes(content)
        try:
            cached = self.cached_parse(digest) if self.cached_parse else None
            chunks, gaps = cached if cached is not None else parse_file(path, self.check, max(.1, self.timeout()), self.recognizer)
        except Exception as exc:
            from .budget import BudgetExceeded
            if isinstance(exc, BudgetExceeded):
                raise
            chunks, gaps = [], ["解析失败：" + type(exc).__name__]
        for i, chunk in enumerate(chunks, 1):
            chunk.update(ref=f"{identifier}:A{i}", line=i)
        date_basis = None
        if not disclosure_date and url:
            match = re.search(r"(?<!\d)(20\d{2})[/-]?(\d{2})[/-]?(\d{2})(?!\d)", url)
            if match:
                candidate = "-".join(match.groups())
                try:
                    cutoff_time(candidate)
                    disclosure_date, date_basis = candidate, "dated_source_url"
                except ValueError:
                    pass
        if not disclosure_date and suffix in {".html", ".htm"}:
            from bs4 import BeautifulSoup
            soup = BeautifulSoup(content, "html.parser")
            meta = soup.find("meta", attrs={"property": "article:published_time"})
            if meta and meta.get("content"):
                try:
                    cutoff_time(meta["content"])
                    disclosure_date, date_basis = meta["content"], "article_published_time"
                except ValueError:
                    pass
        if not disclosure_date:
            for chunk in chunks:
                match = re.search(r"(20\d{2}-\d{2}-\d{2}).{0,12}(?:披露|published)|(?:披露|published).{0,12}(20\d{2}-\d{2}-\d{2})", chunk["text"], re.I)
                if match:
                    disclosure_date = next(x for x in match.groups() if x)
                    date_basis = chunk["ref"]
                    break
        entry = {"id": identifier, "file": path.name, "source_hash": digest, "origin": origin, "url": url,
                 "acquired_at": now(), "date": disclosure_date, "date_basis": date_basis or (channel or {}).get("publication_basis"), "chunks": chunks, "gaps": gaps,
                 "acquisitions": [channel or {"origin": origin, "url": url, "acquired_at": now()}]}
        index["materials"].append(entry)
        self._write(index)
        return entry

    def fetch_raw(self, url):
        self.check()
        parsed = urlparse(url)
        if parsed.scheme not in {"http", "https"} or parsed.username or parsed.password:
            raise ValueError("public HTTP(S) URL required")
        try:
            with httpx.Client(follow_redirects=True, timeout=max(.1, min(45, self.timeout())), headers={"User-Agent": "Yantra/0.3 public research"}) as client:
                with client.stream("GET", url) as response:
                    response.raise_for_status()
                    content = bytearray()
                    for chunk in response.iter_bytes():
                        self.check()
                        content.extend(chunk)
                        if len(content) > 40_000_000:
                            raise ValueError("material exceeds 40 MB limit")
                    mime = response.headers.get("content-type", "")
                    final_url = str(response.url)
        except (httpx.HTTPStatusError, httpx.TransportError):
            # Some disclosure sites reject Python's TLS/client fingerprint.
            # DDGS already depends on primp; use its browser-compatible HTTP client once.
            import primp
            self.check()
            response = primp.Client(impersonate="chrome", timeout=max(.1, min(45, self.timeout()))).get(url)
            if response.status_code >= 400:
                raise OSError(f"public source HTTP {response.status_code}")
            content = response.content
            if len(content) > 40_000_000:
                raise ValueError("material exceeds 40 MB limit")
            mime = response.headers.get("content-type", "")
            final_url = url
        suffix = Path(urlparse(final_url).path).suffix.lower()
        if bytes(content).startswith(b"%PDF"):
            suffix = ".pdf"
        elif "html" in mime or suffix not in FORMATS:
            suffix = ".html" if "html" in mime else ".txt"
        return {"content": bytes(content), "suffix": suffix, "mime": mime, "final_url": final_url}

    def read_materials(self, arguments: ReadArguments):
        index = self.load_manifest()
        documents, gaps = [], list(index["gaps"])
        for entry in index["materials"]:
            if arguments.material_ids is not None and entry["id"] not in arguments.material_ids:
                continue
            if arguments.cutoff and entry["date"] and cutoff_time(entry["date"]) > cutoff_time(arguments.cutoff):
                continue
            if arguments.cutoff and not entry["date"] and cutoff_time(arguments.cutoff).date() < cutoff_time(entry["acquired_at"]).date():
                gaps.append(f"历史截止过滤：未知披露日期的材料已排除 {entry['id']}")
                continue
            content = (self.materials_dir / entry["file"]).read_bytes()
            if hashlib.sha256(content).hexdigest() != entry["source_hash"]:
                raise ValueError("original material hash changed")
            documents.append({**entry, "file_path": str((self.materials_dir / entry["file"]).resolve()),
                              "lines": entry["chunks"], "text": "\n".join(x["text"] for x in entry["chunks"])})
            gaps.extend(entry["gaps"])
        return {"company": index["company"], "company_name": index["company_name"], "material_version": index["version"],
                "fictional": False, "adapter": "raw-v2", "cutoff": arguments.cutoff, "documents": documents, "gaps": list(dict.fromkeys(gaps))}

    def pipeline(self):
        from .sources import SourcePipeline
        return SourcePipeline(self, getattr(self, "source_adapter", None), getattr(self, "checkpoint_store", None), getattr(self, "cutoff", None))

    def fetch(self, url):
        from .sources import Candidate
        receipt = self.pipeline().acquire(Candidate("web_search", "public_web", url=url))
        if not receipt.get("source_id"):
            raise ValueError("source acquisition failed: " + str(receipt.get("failure_reason")))
        return next(e for e in self.load_manifest()["materials"] if e["id"] == receipt["source_id"])

    def search(self, query, purpose, max_results):
        receipt = self.pipeline().discover_and_fetch(query, purpose, max_results)
        with (self.materials_dir / "searches.jsonl").open("a") as file:
            file.write(json.dumps(receipt, ensure_ascii=False) + "\n")
        return receipt
